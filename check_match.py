import os
import glob
import sys
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

def find_files():
    # Find all CSV files
    all_csvs = [f for f in glob.glob("*.csv") if not os.path.basename(f).startswith("~$")]
    
    bcard_csv = None
    item_select_csv = None
    
    for f in all_csvs:
        fname = os.path.basename(f)
        if "名刺" in fname or "Business Card" in fname:
            bcard_csv = f
        elif "項目選択" in fname or "Item" in fname or "item" in fname:
            item_select_csv = f

    # Fallback if names differ
    if not bcard_csv and all_csvs:
        bcard_csv = all_csvs[0]
    
    # Find Excel file
    xlsx_files = [f for f in glob.glob("*.xlsx") if not os.path.basename(f).startswith("~$") and "Report" not in f]
    xlsx_path = xlsx_files[0] if xlsx_files else None

    return bcard_csv, item_select_csv, xlsx_path

def run_comparison(bcard_csv=None, item_select_csv=None, excel_file=None, output_excel="Match_Comparison_Report.xlsx"):
    auto_bcard, auto_item, auto_xlsx = find_files()
    bcard_csv = bcard_csv or auto_bcard
    item_select_csv = item_select_csv or auto_item
    excel_file = excel_file or auto_xlsx

    if not excel_file or not os.path.exists(excel_file):
        print(f"[ERROR] Excel file not found!")
        return

    print("=" * 75)
    print("SanSan Operation Records vs CIC Daily Capacity Comparison Tool")
    print("=" * 75)
    print(f"Excel Target File : {os.path.basename(excel_file)}")
    if bcard_csv and os.path.exists(bcard_csv):
        print(f"Business Card CSV : {os.path.basename(bcard_csv)}")
    if item_select_csv and os.path.exists(item_select_csv):
        print(f"Item Select CSV   : {os.path.basename(item_select_csv)}")
    print("-" * 75)

    # 1. Configuration of all 12 target fields in Excel
    FIELD_CONFIG = {
        'Email':         {'excel_name': 'Email',         'col_idx': 7},
        'URL':           {'excel_name': 'URL',           'col_idx': 8},
        'Date':          {'excel_name': 'Date',          'col_idx': 9},
        'Fax':           {'excel_name': 'Fax',           'col_idx': 10},
        'Full Name':     {'excel_name': 'Full Name',     'col_idx': 11},
        'Item Select':   {'excel_name': 'Item Select',   'col_idx': 12},
        'Company Name':  {'excel_name': 'Company Name',  'col_idx': 13},
        'Division Name': {'excel_name': 'Division Name', 'col_idx': 14},
        'Position Name': {'excel_name': 'Position Name', 'col_idx': 15},
        'Address':       {'excel_name': 'Address',       'col_idx': 16},
        'TEL':           {'excel_name': 'TEL',           'col_idx': 17},
        'MOBILE':        {'excel_name': 'MOBILE',        'col_idx': 18},
    }

    # Mapping from Business Card CSV temp column to FIELD_CONFIG key
    BCARD_MAP = {
        'Date': 'Date',
        'Email': 'Email',
        'FAX': 'Fax',
        'full name': 'Full Name',
        'mobile': 'MOBILE',
        'TEL': 'TEL',
        'URL': 'URL',
        'Address': 'Address',
        'company': 'Company Name',
        'division': 'Division Name',
        'position': 'Position Name'
    }

    csv_records = {} # (uid, field_key) -> matchRate
    csv_users = {}   # uid -> name

    # Load Business Card CSV
    if bcard_csv and os.path.exists(bcard_csv):
        df_bc = pd.read_csv(bcard_csv)
        for _, row in df_bc.iterrows():
            uid = str(row['userIdentifier']).strip()
            temp_val = str(row['temp']).strip() if pd.notna(row.get('temp')) else ''
            name = str(row['name']).strip() if ('name' in row and pd.notna(row['name'])) else ''
            if name:
                csv_users[uid] = name
            
            if temp_val in BCARD_MAP:
                field_key = BCARD_MAP[temp_val]
                try:
                    rate = round(float(row['matchRate']), 2)
                    csv_records[(uid, field_key)] = rate
                except (ValueError, TypeError):
                    csv_records[(uid, field_key)] = str(row['matchRate']).strip()

    # Load Item Select CSV
    if item_select_csv and os.path.exists(item_select_csv):
        df_item = pd.read_csv(item_select_csv)
        for _, row in df_item.iterrows():
            uid = str(row['userIdentifier']).strip()
            if uid not in csv_users:
                csv_users[uid] = ''
            try:
                rate = round(float(row['matchRate']), 2)
                csv_records[(uid, 'Item Select')] = rate
            except (ValueError, TypeError):
                csv_records[(uid, 'Item Select')] = str(row['matchRate']).strip()

    # 2. Load Excel
    wb = openpyxl.load_workbook(excel_file, data_only=True)
    sheet_name = 'Monthly Record' if 'Monthly Record' in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    # Map userIdentifier -> (row_number, operator_name)
    excel_users = {}
    for r in range(4, ws.max_row + 1):
        uid_cell = ws.cell(row=r, column=2).value
        if uid_cell is not None and str(uid_cell).strip() != '':
            uid = str(uid_cell).strip()
            op_name = ws.cell(row=r, column=3).value or ''
            excel_users[uid] = (r, str(op_name).strip())

    # 3. Perform Comparison
    all_uids = sorted(list(set(list(csv_users.keys()) + list(excel_users.keys()))))
    
    results = []
    field_stats = {k: {'checked': 0, 'match': 0, 'mismatch': 0} for k in FIELD_CONFIG}
    mismatch_count = 0
    missing_in_excel_count = 0
    missing_in_csv_count = 0
    match_count = 0

    for uid in all_uids:
        in_excel = uid in excel_users
        row_num, op_name = excel_users.get(uid, (None, csv_users.get(uid, '')))

        for field_key, config in FIELD_CONFIG.items():
            excel_col_name = config['excel_name']
            col_idx = config['col_idx']

            c_val = csv_records.get((uid, field_key))

            e_val = None
            if in_excel and row_num:
                raw_e = ws.cell(row=row_num, column=col_idx).value
                if raw_e is not None and str(raw_e).strip() != '':
                    try:
                        e_val = round(float(raw_e), 2)
                    except (ValueError, TypeError):
                        e_val = str(raw_e).strip()

            # If both are empty, ignore
            if c_val is None and e_val is None:
                continue

            field_stats[field_key]['checked'] += 1
            status = "MATCH"
            diff = 0
            if c_val is not None and e_val is not None:
                if c_val == e_val:
                    status = "MATCH"
                    match_count += 1
                    field_stats[field_key]['match'] += 1
                else:
                    status = "MISMATCH"
                    mismatch_count += 1
                    field_stats[field_key]['mismatch'] += 1
                    try:
                        diff = round(c_val - e_val, 2)
                    except:
                        diff = "N/A"
            elif c_val is not None and e_val is None:
                status = "MISSING_IN_EXCEL"
                missing_in_excel_count += 1
                field_stats[field_key]['mismatch'] += 1
            elif c_val is None and e_val is not None:
                status = "MISSING_IN_CSV"
                missing_in_csv_count += 1
                field_stats[field_key]['mismatch'] += 1

            results.append({
                'User Identifier': uid,
                'Operator Name': op_name,
                'Excel Row': row_num if row_num else 'Not Found',
                'Field': excel_col_name,
                'CSV matchRate': c_val if c_val is not None else '(empty)',
                'Excel Value': e_val if e_val is not None else '(empty)',
                'Difference (CSV - Excel)': diff,
                'Status': status
            })

    # Summary Display
    print("\n--- FIELD BY FIELD BREAKDOWN ---")
    print(f"{'Field Name':<16} | {'Checked':<8} | {'Matched':<8} | {'Mismatch':<8}")
    print("-" * 50)
    for f_key in ['Item Select', 'Date', 'Email', 'Fax', 'Full Name', 'MOBILE', 'TEL', 'URL', 'Address', 'Company Name', 'Division Name', 'Position Name']:
        stats = field_stats[f_key]
        if stats['checked'] > 0:
            print(f"{f_key:<16} | {stats['checked']:<8} | {stats['match']:<8} | {stats['mismatch']:<8}")
    
    print("-" * 75)
    print("--- OVERALL SUMMARY ---")
    print(f"Total Values Checked: {len(results)}")
    print(f"Matched Values      : {match_count}")
    print(f"Mismatches          : {mismatch_count}")
    print(f"Missing in Excel    : {missing_in_excel_count}")
    print(f"Missing in CSV      : {missing_in_csv_count}")
    print("-" * 75)

    if mismatch_count == 0 and missing_in_excel_count == 0 and missing_in_csv_count == 0:
        print("[SUCCESS] PERFECT MATCH! All values in CSV files match the Excel file exactly.")
    else:
        print(f"[WARNING] Discrepancies detected!")
        for r in results:
            if r['Status'] != 'MATCH':
                print(f"  - [{r['Status']}] User: {r['User Identifier']} | Field: {r['Field']} | CSV: {r['CSV matchRate']} vs Excel: {r['Excel Value']}")

    # 4. Generate Detailed Excel Report
    df_out = pd.DataFrame(results)

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Comparison Summary"

    headers = list(df_out.columns)
    ws_out.append(headers)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for col_num, h in enumerate(headers, 1):
        cell = ws_out.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    match_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    mismatch_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    for row_data in df_out.values.tolist():
        ws_out.append(row_data)
        curr_row = ws_out.max_row
        status = row_data[-1]
        row_fill = match_fill if status == "MATCH" else mismatch_fill
        for c in range(1, len(row_data) + 1):
            ws_out.cell(row=curr_row, column=c).fill = row_fill

    for col in ws_out.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws_out.column_dimensions[col_letter].width = max(max_len + 4, 12)

    report_path = os.path.join(os.path.dirname(os.path.abspath(excel_file)), output_excel)
    wb_out.save(report_path)
    print(f"\nDetailed report saved to: {report_path}")
    print("=" * 75)

if __name__ == "__main__":
    bcard_arg = sys.argv[1] if len(sys.argv) > 1 else None
    item_arg = sys.argv[2] if len(sys.argv) > 2 else None
    xlsx_arg = sys.argv[3] if len(sys.argv) > 3 else None
    run_comparison(bcard_arg, item_arg, xlsx_arg)
